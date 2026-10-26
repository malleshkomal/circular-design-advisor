"""Build outputs/lci_workbook.html: an offline, browsable view of outputs/LCI_inventory.xlsx.

Shows every sheet with its calculated values; click a calculated cell to see its Excel formula.
The workbook must hold calculated values: after build_lci_workbook.py, open and save it in
Excel or LibreOffice (or recalculate it headless) before running this script.

Usage:  python scripts/build_lci_viewer.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "outputs" / "LCI_inventory.xlsx"
TEMPLATE = ROOT / "web" / "lci_viewer_template.html"
OUT = ROOT / "outputs" / "lci_workbook.html"

FILL_CLASS = {"1F4E3D": "h", "D9EAD3": "s", "FFFF00": "k", "EDEDED": "t"}


def fmt(v, nf):
    """Render a cell value roughly as Excel displays it with its number format."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (int, float)):
        if "%" in nf:
            return "–" if abs(v) < 1e-12 else f"{v * 100:.1f}%"
        if "€" in nf:
            return "–" if abs(v) < 1e-12 else (f"€{v:,.2f}" if v >= 0 else f"(€{-v:,.2f})")
        if nf == "0":
            return f"{v:.0f}"
        if abs(v) < 1e-12 and "-" in nf:
            return "–"
        if isinstance(v, int) or float(v).is_integer():
            return f"{v:,.0f}" if abs(v) >= 1000 else f"{v:g}"
        d = 6 if "000000" in nf else 4 if "0000" in nf else 3 if "0.000" in nf else 2 if "0.00" in nf else 3
        s = f"{abs(v):,.{d}f}"
        return f"({s})" if v < 0 else s
    return str(v)


def sheets() -> list[dict]:
    values = load_workbook(XLSX, data_only=True)
    formulas = load_workbook(XLSX)
    out = []
    for ws in values.worksheets:
        wf = formulas[ws.title]
        rows = []
        for r in range(1, ws.max_row + 1):
            cells = []
            for c in range(1, ws.max_column + 1):
                cell, f = ws.cell(r, c), wf.cell(r, c).value
                is_formula = isinstance(f, str) and f.startswith("=")
                if is_formula and cell.value is None:
                    raise SystemExit(f"{XLSX.name} has no calculated values ({ws.title}!{cell.coordinate}): "
                                     "open and save it in Excel or LibreOffice first.")
                cls = []
                if cell.fill is not None and cell.fill.fill_type:
                    rgb = cell.fill.fgColor.rgb if isinstance(cell.fill.fgColor.rgb, str) else ""
                    cls += [v for k, v in FILL_CLASS.items() if rgb.endswith(k)]
                color = cell.font.color.rgb if cell.font and cell.font.color and isinstance(cell.font.color.rgb, str) else ""
                if color.endswith("0000FF"):
                    cls.append("in")
                if color.endswith("C00000"):
                    cls.append("as")
                if cell.font and cell.font.b:
                    cls.append("b")
                if is_formula:
                    cls.append("f")
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cls.append("n")
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cls.append("u")
                cells.append([fmt(cell.value, cell.number_format), " ".join(cls), f if is_formula else ""])
            while cells and cells[-1][0] == "" and not cells[-1][1]:
                cells.pop()
            rows.append(cells)
        out.append({"name": ws.title, "rows": rows})
    return out


def render() -> str:
    page = TEMPLATE.read_text(encoding="utf-8").split("-->", 1)[1].lstrip()
    page = page.replace("__DATA__", json.dumps(sheets(), ensure_ascii=False))
    head, body = page.split('<div class="wrap">', 1)
    return ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
            f'{head}</head>\n<body>\n<div class="wrap">{body}</body>\n</html>\n')


def build() -> Path:
    OUT.write_text(render(), encoding="utf-8")
    return OUT


if __name__ == "__main__":
    print(build())
