"""Build outputs/LCI_inventory.xlsx: a formula-driven life cycle inventory (LCI) of one aluminium booth panel.

The workbook mirrors src/cirdesign/model.py in Excel formulas:
    Parameters + Context  ->  Calc (per strategy)  ->  LCI (flows)  ->  LCIA_GWP, Cost  ->  KPIs
Every input lives in one cell (blue text); everything else is a formula, so changing a parameter,
the design, the business model or the credit method recalculates the whole inventory.
The Validation sheet compares the Excel KPIs with the Python model at the default context.

Usage:  python scripts/build_lci_workbook.py   (then recalculate in Excel or LibreOffice)
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from cirdesign.data import load_all  # noqa: E402
from cirdesign.model import DEFAULT_CONTEXT, evaluate_all  # noqa: E402
from cirdesign.recommend import WEIGHTS, score_strategies  # noqa: E402

OUT = ROOT / "outputs" / "LCI_inventory.xlsx"

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F4E3D")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="1F4E3D")
FILL_SECTION = PatternFill("solid", fgColor="D9EAD3")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_TOTAL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

STRATS = ["reuse", "repair", "reprocessing", "recycling", "disposal"]
SCOLS = ["D", "E", "F", "G", "H"]  # one column per strategy on every strategy sheet
HEADER_ROW = 4

NUM = "#,##0.000;(#,##0.000);-"
NUM2 = "#,##0.00;(#,##0.00);-"
PCT = "0.0%;(0.0%);-"
EUR = "€#,##0.00;(€#,##0.00);-"


def style_range(ws, ref, font=None, fill=None, fmt=None, border=None, align=None):
    for row in ws[ref]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if fmt:
                c.number_format = fmt
            if border:
                c.border = border
            if align:
                c.alignment = align


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_NOTE


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


class StrategySheet:
    """A sheet whose rows hold one formula per strategy in columns D:H.

    Formula templates use {key} for a row on the same sheet and {Calc:key} for a row on Calc;
    both resolve to the same strategy column.
    """

    registry: dict[str, dict[str, int]] = {}

    def __init__(self, wb, name, title_text, sub, label_head="Item"):
        self.ws = wb.create_sheet(name)
        self.name = name
        self.rows: dict[str, int] = {}
        StrategySheet.registry[name] = self.rows
        title(self.ws, title_text, sub)
        header(self.ws, HEADER_ROW, [label_head, "Unit", "Formula / note"] + STRATS + ["Basis / source"])
        for col, i in zip(SCOLS, range(len(STRATS))):
            c = self.ws[f"{col}{HEADER_ROW}"]
            c.value = f"=Strategies!A{i + 2}"
        self.r = HEADER_ROW + 1
        widths(self.ws, {"A": 46, "B": 11, "C": 44, **{c: 14 for c in SCOLS}, "I": 58})
        self.ws.freeze_panes = f"D{HEADER_ROW + 1}"

    def resolve(self, tpl, col):
        def sub(m):
            ref = m.group(1)
            if ":" in ref:
                sheet, key = ref.split(":")
                return f"{sheet}!{col}{StrategySheet.registry[sheet][key]}"
            return f"{col}{self.rows[ref]}"
        return "=" + re.sub(r"\{([A-Za-z_:]+)\}", sub, tpl)

    def section(self, text):
        self.ws.cell(row=self.r, column=1, value=text)
        style_range(self.ws, f"A{self.r}:I{self.r}", font=F_BOLD, fill=FILL_SECTION)
        self.r += 1

    def row(self, key, label, unit, tpl, note="", basis="", fmt=NUM, bold=False, fill=None):
        r = self.r
        self.rows[key] = r
        ws = self.ws
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=unit)
        ws.cell(row=r, column=3, value=note)
        ws.cell(row=r, column=9, value=basis)
        for col in SCOLS:
            ws[f"{col}{r}"] = self.resolve(tpl, col)
        style_range(ws, f"A{r}:I{r}", font=F_BOLD if bold else F_BASE, border=BOX)
        style_range(ws, f"D{r}:H{r}", fmt=fmt)
        style_range(ws, f"C{r}:C{r}", font=F_NOTE)
        style_range(ws, f"I{r}:I{r}", font=F_NOTE)
        if fill:
            style_range(ws, f"A{r}:H{r}", fill=fill)
        self.r += 1
        return r

    def note(self, text):
        self.r += 1
        self.ws.cell(row=self.r, column=1, value=text).font = F_NOTE
        self.r += 1


def build():
    T = load_all()
    src, designs, bms, strats = T["sources"], T["designs"], T["business_models"], T["strategies"]
    wb = Workbook()

    # ------------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    title(ws, "Circular Design Advisor – Life Cycle Inventory (LCI)",
          "One modular aluminium exhibition-booth panel · five end-of-life / circular strategies · formula-driven")
    lines = [
        ("Goal and scope", None),
        ("Functional unit", "ONE USE of one panel at one exhibition event (per use). Lifetime burdens are divided by the expected number of uses."),
        ("Reference flow", "1 panel of the selected design (mass from Product_Design), over its expected service life."),
        ("System boundary", "Cradle-to-grave: panel production (metal + fabrication), transport to/from events, cleaning, repair, reprocessing, end-of-life transport, sorting, waste-to-energy; recycling credit (net scrap) or cut-off."),
        ("Impact category", "Climate change (GWP100, kg CO2e) only. Costs are a proxy in EUR."),
        ("Multifunctionality", "Net-scrap method (default): credit only for recovered metal beyond the recycled input share, valued at (EF primary − EF secondary). Cut-off: no credit. Switch on the Context sheet."),
        ("Geography / time", "Europe / Germany, 2023–2025 data (see Parameters)."),
        ("", None),
        ("Workbook map", None),
        ("Parameters", "All 30 model parameters with value, range, unit, data type, quality score (1–5), source and URL. Edit column C (blue) to change the model."),
        ("Product_Design / Strategies / Business_Models", "Design archetypes, strategy settings and business-model behaviour (all flagged assumptions)."),
        ("Context", "Decision context: design, business model, planned uses, event and end-of-life distances, credit method (yellow cells). Also derived global factors."),
        ("Calc", "Per-strategy intermediate quantities: planned uses, loss probability, expected uses, repair events, recovery yield, credit mass."),
        ("LCI", "THE INVENTORY: every foreground flow (kg, kWh, t·km) per panel life and per functional unit, with an indicative background dataset for Brightway/ecoinvent mapping."),
        ("LCIA_GWP", "LCI per functional unit × emission factor = kg CO2e per use, by flow and by life-cycle stage."),
        ("Cost", "Cost inventory per panel life and per use (EUR, incl. business-model margin)."),
        ("KPIs", "Climate, cost, circularity (simplified MCI), material recovery, virgin Al, feasibility, weighted score, rank and recommendation."),
        ("Validation", "Excel KPIs vs Python model (src/cirdesign/model.py) at the default context. All rows should read OK."),
        ("", None),
        ("Colour legend", None),
        ("Blue text", "Hard-coded input – edit these."),
        ("Yellow fill", "Key decision levers (Context sheet)."),
        ("Green text", "Link to another sheet."),
        ("Black text", "Formula – do not overwrite."),
        ("", None),
        ("Important", None),
        ("Data status", "Only 6 of the model inputs come from literature or public databases; the rest are author assumptions (type = assumption, quality 2). Results illustrate the method, not a real product. Do not cite the numbers."),
        ("Regenerate", "python scripts/build_lci_workbook.py  (rebuilds this file from data/*.csv; open in Excel to recalculate)."),
    ]
    r = 4
    for k, v in lines:
        a = ws.cell(row=r, column=1, value=k)
        if v is None and k:
            a.font, a.fill = F_HEAD, FILL_HEAD
            ws.cell(row=r, column=2).fill = FILL_HEAD
        else:
            a.font = F_BOLD
            b = ws.cell(row=r, column=2, value=v)
            b.font, b.alignment = F_BASE, WRAP
        if k == "Blue text":
            a.font = Font(name=FONT, size=10, bold=True, color="0000FF")
        elif k == "Yellow fill":
            a.fill = FILL_KEY
        elif k == "Green text":
            a.font = Font(name=FONT, size=10, bold=True, color="008000")
        r += 1
    widths(ws, {"A": 42, "B": 120})

    # ------------------------------------------------------------------ Parameters
    ws = wb.create_sheet("Parameters")
    title(ws, "Parameters", "Source: data/literature_sources.csv. Column C is the value used by the model; edit it to run a scenario. Low/High = literature or assumed range.")
    cols = ["ID (named range)", "Parameter", "Value", "Low", "High", "Unit", "Type", "Used in model", "Quality (1–5)", "Source", "Year", "URL", "Note"]
    header(ws, 4, cols)
    for i, row in enumerate(src.itertuples(), start=5):
        vals = [row.id, row.parameter, row.value, row.low, row.high, row.unit, row.type, bool(row.used_in_model),
                row.quality, row.source, row.year, row.url, row.note]
        for j, v in enumerate(vals, start=1):
            if isinstance(v, float) and v != v:  # NaN
                v = None
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = F_BASE, BOX
        ws.cell(row=i, column=3).font = F_INPUT
        for j in (4, 5):
            ws.cell(row=i, column=j).font = F_INPUT
        if row.type == "assumption":
            ws.cell(row=i, column=7).font = Font(name=FONT, size=10, color="C00000")
        add_name(wb, row.id, f"Parameters!$C${i}")
    last = 4 + len(src)
    style_range(ws, f"C5:E{last}", fmt="#,##0.000")
    style_range(ws, f"J5:M{last}", align=WRAP)
    widths(ws, {"A": 26, "B": 46, "C": 11, "D": 10, "E": 10, "F": 16, "G": 12, "H": 10, "I": 10,
                "J": 48, "K": 9, "L": 40, "M": 60})
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:M{last}"

    # ------------------------------------------------------------------ lookup tables
    def table_sheet(name, sub, df, index_name, input_cols):
        ws = wb.create_sheet(name)
        df = df.reset_index()
        heads = list(df.columns)
        ws.cell(row=1, column=1)  # tables start on row 1 so INDEX/MATCH can use whole columns
        header(ws, 1, heads)
        for i, rr in enumerate(df.itertuples(index=False), start=2):
            for j, v in enumerate(rr, start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = BOX
                c.font = F_INPUT if heads[j - 1] in input_cols else F_BASE
        n = len(df) + 1
        ws.cell(row=n + 2, column=1, value=sub).font = F_NOTE
        widths(ws, {get_column_letter(j): (60 if h == "note" else 16) for j, h in enumerate(heads, start=1)})
        ws.freeze_panes = "B2"
        return ws, heads, n

    _, dheads, dn = table_sheet("Product_Design", "Source: data/product_data.csv. Design archetypes are assumptions.",
                                designs, "design", [c for c in designs.columns if c not in ("name", "note")])
    _, sheads, sn = table_sheet("Strategies", "Source: data/strategies.csv. repair / reprocess are 0/1 switches. Row order fixes the strategy columns on every other sheet.",
                                strats, "strategy", [c for c in strats.columns if c not in ("label", "family", "note")])
    _, bheads, bn = table_sheet("Business_Models", "Source: data/business_models.csv (all assumed).",
                                bms, "name", ["redeploy", "margin", "damage_mult"])
    assert list(strats.index) == STRATS

    # ------------------------------------------------------------------ Context
    ws = wb.create_sheet("Context")
    title(ws, "Decision context and derived global factors", "Yellow cells are the levers. Defaults reproduce the README baseline (design B, rental, 6 planned uses, 400 km).")
    header(ws, 4, ["Setting", "Value", "Unit", "Note"])
    ctx = [
        ("ctx_design", "Panel design", DEFAULT_CONTEXT["design"], "", "A, B or C (Product_Design)"),
        ("ctx_bm", "Business model", DEFAULT_CONTEXT["business_model"], "", "sale, rental or paas (Business_Models)"),
        ("ctx_n", "Planned uses (reuse strategies)", DEFAULT_CONTEXT["n_uses"], "events", "Capped at the design's technical lifetime"),
        ("ctx_event_km", "Distance to each event (one way)", DEFAULT_CONTEXT["event_km"], "km", "Reuse routes travel out and back each event"),
        ("ctx_eol_km", "Distance to end-of-life facility", DEFAULT_CONTEXT["eol_km"], "km", ""),
        ("ctx_credit", "Recycling credit method", DEFAULT_CONTEXT["credit"], "", "net_scrap or cut_off"),
    ]
    for i, (nm, lab, v, unit, note) in enumerate(ctx, start=5):
        ws.cell(row=i, column=1, value=lab).font = F_BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.fill, c.border = F_INPUT, FILL_KEY, BOX
        ws.cell(row=i, column=3, value=unit).font = F_BASE
        ws.cell(row=i, column=4, value=note).font = F_NOTE
        add_name(wb, nm, f"Context!$B${i}")
    for ref, opts in (("B5", ",".join(designs.index)), ("B6", ",".join(bms.index)), ("B10", "net_scrap,cut_off")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(ref)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="100")
    ws.add_data_validation(dv)
    dv.add("B7")

    ws["A12"] = "Selected design and business-model attributes"
    style_range(ws, "A12:D12", font=F_BOLD, fill=FILL_SECTION)
    dcol = {h: get_column_letter(j) for j, h in enumerate(dheads, start=1)}
    bcol = {h: get_column_letter(j) for j, h in enumerate(bheads, start=1)}
    attrs = [
        ("d_mass", "Panel mass", "mass_kg", "kg", "D"),
        ("d_repair", "Repairability", "repairability", "share", "D"),
        ("d_fab", "Fabrication energy factor", "fab_factor", "×", "D"),
        ("d_cost", "Manufacturing cost factor", "cost_factor", "×", "D"),
        ("d_dmg", "Design damage multiplier", "damage_mult", "×", "D"),
        ("d_life", "Technical lifetime", "technical_lifetime_events", "events", "D"),
        ("d_disasm", "Disassembly score", "disassembly", "share", "D"),
        ("bm_redeploy", "Redeployment probability per event", "redeploy", "share", "B"),
        ("bm_margin", "Business-model margin", "margin", "share", "B"),
        ("bm_dmg", "Business-model damage multiplier", "damage_mult", "×", "B"),
    ]
    r = 13
    for nm, lab, field, unit, kind in attrs:
        ws.cell(row=r, column=1, value=lab).font = F_BASE
        if kind == "D":
            f = f"=INDEX(Product_Design!${dcol[field]}$2:${dcol[field]}${dn},MATCH(ctx_design,Product_Design!$A$2:$A${dn},0))"
        else:
            f = f"=INDEX(Business_Models!${bcol[field]}$2:${bcol[field]}${bn},MATCH(ctx_bm,Business_Models!$A$2:$A${bn},0))"
        c = ws.cell(row=r, column=2, value=f)
        c.font, c.border, c.number_format = F_LINK, BOX, "#,##0.00"
        ws.cell(row=r, column=3, value=unit).font = F_BASE
        add_name(wb, nm, f"Context!$B${r}")
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Derived global factors")
    style_range(ws, f"A{r}:D{r}", font=F_BOLD, fill=FILL_SECTION)
    r += 1
    derived = [
        ("ef_mix", "Aluminium input mix factor", "=(1-recycled_content)*ef_primary+recycled_content*ef_secondary", "kg CO2e/kg", "Recycled-content weighted mix of primary and secondary Al"),
        ("leg_ef", "Effective freight factor", "=transport_ef/transport_load_factor", "kg CO2e/(t·km)", "Full-load factor divided by load factor"),
        ("price_kg", "Aluminium price", "=lme_price_usd_t*usd_eur/1000", "EUR/kg", "LME price converted to EUR"),
        ("recov_yield", "Recovery yield, recycling route", "=collection_rate*melt_yield", "share", "Collection × remelting yield"),
        ("delta_ef", "Credit factor (avoided primary)", "=ef_primary-ef_secondary", "kg CO2e/kg", "Burden avoided per kg of secondary Al that replaces primary"),
        ("dist_factor", "Distance penalty (feasibility)", "=MAX(0,1-ctx_event_km/2000)", "share", "Linear penalty, zero at 2,000 km"),
    ]
    for nm, lab, f, unit, note in derived:
        ws.cell(row=r, column=1, value=lab).font = F_BASE
        c = ws.cell(row=r, column=2, value=f)
        c.font, c.border, c.number_format = F_BASE, BOX, "#,##0.0000"
        ws.cell(row=r, column=3, value=unit).font = F_BASE
        ws.cell(row=r, column=4, value=note).font = F_NOTE
        add_name(wb, nm, f"Context!$B${r}")
        r += 1
    widths(ws, {"A": 40, "B": 14, "C": 16, "D": 62})

    # ------------------------------------------------------------------ Calc
    S = StrategySheet(wb, "Calc", "Per-strategy intermediate quantities",
                      "Life-cycle logic per strategy (mirrors model.evaluate). Columns D:H follow the Strategies sheet order.", "Quantity")
    S.section("Strategy settings (from Strategies sheet)")
    sc = {h: get_column_letter(j) for j, h in enumerate(sheads, start=1)}

    def slook(field):
        return f"INDEX(Strategies!${sc[field]}$2:${sc[field]}${sn},MATCH({{strat}},Strategies!$A$2:$A${sn},0))"

    S.rows["strat"] = HEADER_ROW
    S.row("single", "Single-use route (1 = yes)", "flag", f'IF({slook("family")}="single",1,0)', "family = single", fmt="0")
    S.row("st_repair", "Repair switch", "flag", slook("repair"), fmt="0")
    S.row("st_reprocess", "Reprocessing switch", "flag", slook("reprocess"), fmt="0")
    S.row("st_dmgf", "Strategy damage factor", "×", slook("damage_factor"), fmt=NUM2)
    S.row("st_labour", "Workshop labour per event", "h", slook("labour_h"), fmt=NUM2)
    S.row("st_feas", "Base feasibility", "share", slook("feasibility_base"), fmt=NUM2)
    S.row("is_disposal", "Disposal route (1 = yes)", "flag", 'IF({strat}="disposal",1,0)', fmt="0")
    S.section("Service life")
    S.row("N", "Planned uses N", "events", "IF({single}=1,1,MIN(ROUND(ctx_n,0),d_life))", "min(planned, technical lifetime)", fmt="0")
    S.row("dmg", "Damage probability per event", "share",
          "IF({single}=1,0,damage_rate*d_dmg*bm_dmg*{st_dmgf})", "damage_rate × design × business model × strategy", fmt=PCT)
    S.row("loss_dmg", "Damage that ends the panel's life", "share",
          "IF({st_repair}=1,{dmg}*(1-d_repair),{dmg})", "repairable damage is fixed", fmt=PCT)
    S.row("p_loss", "Loss probability per event", "share",
          "IF({single}=1,1,1-(1-{loss_dmg})*bm_redeploy)", "damage loss or not redeployed", fmt=PCT)
    S.row("U", "Expected uses U", "uses",
          "IF({single}=1,1,IF({p_loss}<=1E-12,MAX({N},1),(1-(1-{p_loss})^MAX({N},1))/{p_loss}))",
          "Σ_{k=0..N-1} (1-p)^k", basis="Geometric survival over N planned events", bold=True)
    S.row("cycles", "Reuse cycles (U − 1)", "cycles", "{U}-1")
    S.row("repair_events", "Repairs per event", "repairs/event",
          "IF(AND({single}=0,{st_repair}=1),{dmg}*d_repair,0)", fmt="0.0000")
    S.row("legs", "Freight legs per event", "legs", "IF({single}=1,1,2)", "reuse: out and back", fmt="0")
    S.section("End of life")
    S.row("y", "Material recovery yield", "share",
          "IF({is_disposal}=1,disposal_recovery,recov_yield)", "bottom-ash recovery for disposal", fmt=PCT)
    S.row("credit_mass", "Credited secondary metal", "kg",
          'IF(ctx_credit="cut_off",0,d_mass*MAX({y}-recycled_content,0))',
          "net scrap: recovered beyond recycled input", basis="Avoids double counting recycled input and EoL credit")

    # ------------------------------------------------------------------ LCI
    L = StrategySheet(wb, "LCI", "Life Cycle Inventory – foreground flows",
                      "Amounts per panel life (rows 'per life') and per functional unit = one use (rows 'per FU'). Background datasets are indicative suggestions for Brightway/ecoinvent mapping – verify before use.",
                      "Flow")
    L.ws["I4"] = "Indicative background dataset (ecoinvent 3.x, verify)"
    flows = [
        # stage, key, label, unit, lifetime formula, EF name, background, note
        ("A. Production (panel manufacture)", None),
        ("A", "f_prim", "Primary aluminium, ingot", "kg", "d_mass*(1-recycled_content)", "ef_primary",
         "market for aluminium, primary, ingot | IAI Area, EU27 & EFTA", "m × (1 − recycled content)"),
        ("A", "f_sec", "Secondary aluminium (post-consumer scrap)", "kg", "d_mass*recycled_content", "ef_secondary",
         "treatment of aluminium scrap, post-consumer, prepared for recycling, at remelter | RER", "m × recycled content"),
        ("A", "f_el_fab", "Electricity, fabrication (cut, machine, finish)", "kWh", "d_mass*e_fab*d_fab", "grid_ef",
         "market for electricity, low voltage | DE", "m × e_fab × design factor"),
        ("B. Use phase (events, maintenance)", None),
        ("B", "f_tkm_ev", "Road freight to and from events", "t·km", "{Calc:U}*{Calc:legs}*d_mass/1000*ctx_event_km", "leg_ef",
         "market for transport, freight, lorry >32 metric ton, EURO6 | RER", "U × legs × m/1000 × km"),
        ("B", "f_el_clean", "Electricity, cleaning and inspection", "kWh", "IF({Calc:single}=1,0,{Calc:U}*d_mass*e_clean)", "grid_ef",
         "market for electricity, low voltage | DE", "reuse routes, every event"),
        ("B", "f_el_repro", "Electricity, reprocessing (strip, re-finish)", "kWh", "{Calc:U}*d_mass*e_reprocess*{Calc:st_reprocess}", "grid_ef",
         "market for electricity, low voltage | DE", "reprocessing route only"),
        ("B", "f_el_rep", "Electricity, repair", "kWh", "{Calc:U}*{Calc:repair_events}*d_mass*e_repair", "grid_ef",
         "market for electricity, low voltage | DE", "repair route only"),
        ("B", "f_rep_prim", "Repair parts, primary aluminium", "kg", "{Calc:U}*{Calc:repair_events}*d_mass*repair_share*(1-recycled_content)", "ef_primary",
         "market for aluminium, primary, ingot | IAI Area, EU27 & EFTA", "replacement share × mass"),
        ("B", "f_rep_sec", "Repair parts, secondary aluminium", "kg", "{Calc:U}*{Calc:repair_events}*d_mass*repair_share*recycled_content", "ef_secondary",
         "treatment of aluminium scrap, post-consumer, prepared for recycling, at remelter | RER", "replacement share × mass"),
        ("C. End of life", None),
        ("C", "f_tkm_eol", "Road freight to end-of-life facility", "t·km", "d_mass/1000*ctx_eol_km", "leg_ef",
         "market for transport, freight, lorry >32 metric ton, EURO6 | RER", "m/1000 × km"),
        ("C", "f_el_sort", "Electricity, dismantling and sorting", "kWh", "IF({Calc:is_disposal}=1,0,d_mass*sort_energy)", "grid_ef",
         "market for electricity, low voltage | DE", "not for disposal"),
        ("C", "f_wte", "Waste-to-energy treatment of aluminium", "kg", "IF({Calc:is_disposal}=1,d_mass,0)", "disposal_ef",
         "treatment of waste aluminium, municipal incineration | CH", "disposal route only"),
        ("D. Avoided burden (credit)", None),
        ("D", "f_avoid", "Avoided primary aluminium (substituted by recovered metal)", "kg", "-{Calc:credit_mass}", "delta_ef",
         "−(primary ingot − secondary) | system expansion", "negative = avoided; 0 under cut-off"),
        ("E. Outputs and memo flows (no direct GWP)", None),
        ("E", "o_recov", "Aluminium recovered to remelting / bottom-ash recovery", "kg", "d_mass*{Calc:y}", None,
         "output to recycling", "m × recovery yield"),
        ("E", "o_lost", "Aluminium lost (slag, residues, uncollected)", "kg", "d_mass*(1-{Calc:y})", None,
         "output to waste", "m × (1 − yield)"),
        ("E", "o_panel", "Panel discarded at end of life", "kg", "d_mass", None, "reference flow mass", "one panel"),
    ]
    L.section("Per panel life (reference flow = 1 panel)")
    flow_meta = []
    for f in flows:
        if f[1] is None:
            L.section(f[0])
            continue
        stage, key, lab, unit, tpl, ef, bg, note = f
        L.row(key, lab, unit, tpl, note, bg)
        flow_meta.append((stage, key, lab, unit, ef, bg))
    L.row("fu", "Functional units delivered (expected uses U)", "uses", "{Calc:U}", "link to Calc", bold=True, fill=FILL_TOTAL)
    for col in SCOLS:
        L.ws[f"{col}{L.rows['fu']}"].font = Font(name=FONT, size=10, bold=True, color="008000")

    L.r += 1
    L.section("Per functional unit (one use of one panel) = per life ÷ U")
    for stage, key, lab, unit, ef, bg in flow_meta:
        L.row(key + "_fu", lab, unit, "{%s}/{fu}" % key, "per life ÷ U", bg)
    L.note("Mass balance check: recovered + lost = panel mass for every strategy (see row below).")
    L.row("mb", "Mass balance (recovered + lost − panel)", "kg", "{o_recov}+{o_lost}-{o_panel}", "should be 0", fmt="0.000000")

    # ------------------------------------------------------------------ LCIA
    G = StrategySheet(wb, "LCIA_GWP", "Impact assessment – climate change (GWP100)",
                      "kg CO2e per functional unit (one use) = LCI per FU × emission factor. EF column links to Parameters / Context.",
                      "Flow")
    G.ws["I4"] = "Emission factor (kg CO2e/unit)"
    G.section("GWP by flow (kg CO2e per use)")
    stage_rows: dict[str, list[str]] = {"A": [], "B": [], "C": [], "D": []}
    for stage, key, lab, unit, ef, bg in flow_meta:
        if ef is None:
            continue
        r = G.row("g_" + key, lab, "kg CO2e", "{LCI:%s_fu}*%s" % (key, ef), f"LCI per FU × {ef}")
        c = G.ws[f"I{r}"]
        c.value, c.font, c.number_format = f"={ef}", F_LINK, "0.0000"
        stage_rows[stage].append("g_" + key)
    G.section("GWP by life-cycle stage (kg CO2e per use)")
    stage_names = {"A": "Panel production", "B": "Use phase (transport, cleaning, repair, reprocessing)",
                   "C": "End-of-life processes", "D": "Recycling credit"}
    for s, keys in stage_rows.items():
        G.row("st_" + s, stage_names[s], "kg CO2e", "+".join("{%s}" % k for k in keys), "sum of stage flows")
    G.row("gwp", "Total climate impact per use", "kg CO2e/use", "{st_A}+{st_B}+{st_C}+{st_D}", bold=True, fill=FILL_TOTAL, fmt=NUM2)
    G.row("gwp_life", "Total climate impact per panel life", "kg CO2e", "{gwp}*{Calc:U}", fmt=NUM2)
    G.section("Stage shares of gross burden (excl. credit)")
    for s in "ABC":
        G.row("sh_" + s, stage_names[s], "share", "IF(({st_A}+{st_B}+{st_C})=0,0,{st_%s}/({st_A}+{st_B}+{st_C}))" % s, fmt=PCT)

    # ------------------------------------------------------------------ Cost
    C = StrategySheet(wb, "Cost", "Cost inventory (EUR proxy)",
                      "Per panel life; negative = revenue. Per-use cost includes the business-model margin.", "Cost item")
    C.section("Per panel life (EUR)")
    C.row("c_panel", "Panel purchase (metal + manufacturing)", "EUR", "d_mass*(price_kg+manuf_cost*d_cost)", "m × (Al price + manuf. cost × factor)", fmt=EUR)
    C.row("c_freight", "Freight to and from events", "EUR", "{Calc:U}*{Calc:legs}*d_mass/1000*ctx_event_km*transport_cost", fmt=EUR)
    C.row("c_handling", "Set-up and dismantling labour", "EUR", "{Calc:U}*handling_cost", fmt=EUR)
    C.row("c_storage", "Storage between events", "EUR", "IF({Calc:single}=1,0,{Calc:U}*storage_cost)", "reuse routes", fmt=EUR)
    C.row("c_labour", "Workshop labour (inspect, repair, reprocess)", "EUR",
          "IF({Calc:single}=1,0,{Calc:U}*({Calc:st_labour}*labour_rate*IF({Calc:st_repair}=1,{Calc:repair_events},1)+IF({Calc:st_repair}=1,0.05*labour_rate,0)))",
          "repair: labour per repair + 0.05 h inspection", fmt=EUR)
    C.row("c_repmat", "Repair material", "EUR",
          "{Calc:U}*{Calc:repair_events}*repair_share*d_mass*(price_kg+manuf_cost*d_cost)", fmt=EUR)
    C.row("c_eolfr", "Freight to end-of-life", "EUR", "d_mass/1000*ctx_eol_km*transport_cost", fmt=EUR)
    C.row("c_scrap", "Scrap revenue", "EUR", "IF({Calc:is_disposal}=1,0,-d_mass*{Calc:y}*price_kg*scrap_price_ratio)", "revenue (negative)", fmt=EUR)
    C.row("c_gate", "Disposal gate fee", "EUR", "IF({Calc:is_disposal}=1,d_mass*gate_fee,0)", fmt=EUR)
    C.row("c_total", "Total cost per panel life", "EUR", "SUM(D%d:D%d)" % (C.rows["c_panel"], C.r - 1), bold=True, fill=FILL_TOTAL, fmt=EUR)
    # SUM template above was written for column D; rewrite per column
    for col in SCOLS:
        C.ws[f"{col}{C.rows['c_total']}"] = f"=SUM({col}{C.rows['c_panel']}:{col}{C.rows['c_gate']})"
    C.row("c_use", "Cost per use (incl. margin)", "EUR/use", "{c_total}/{Calc:U}*(1+bm_margin)", bold=True, fmt=EUR)

    # ------------------------------------------------------------------ KPIs
    K = StrategySheet(wb, "KPIs", "Circularity and sustainability KPIs",
                      "All values per use unless stated. Weighted score = 0.50 environment + 0.25 circularity + 0.15 economics + 0.10 feasibility (edit weights below).",
                      "KPI")
    K.section("Core KPIs")
    K.row("k_U", "Expected uses", "uses", "{Calc:U}", fmt=NUM2)
    K.row("k_cyc", "Reuse cycles", "cycles", "{Calc:cycles}", fmt=NUM2)
    K.row("k_gwp", "Climate impact per use", "kg CO2e/use", "{LCIA_GWP:gwp}", "from LCIA_GWP", bold=True, fmt=NUM2)
    K.row("k_cost", "Cost per use", "EUR/use", "{Cost:c_use}", "from Cost", bold=True, fmt=EUR)
    K.row("k_V", "Virgin feedstock V", "kg", "d_mass*(1-recycled_content)", "MCI input")
    K.row("k_W", "Unrecoverable waste W", "kg", "d_mass*(1-{Calc:y})", "MCI input")
    K.row("k_lfi", "Linear flow index LFI", "–", "({k_V}+{k_W})/(2*d_mass)", "(V + W) / 2M")
    K.row("k_mci", "Circularity (simplified MCI)", "0–1", "MAX(0,1-{k_lfi}*0.9/MAX({Calc:U},1))",
          "1 − LFI × 0.9 / utility; utility = U", basis="Ellen MacArthur Foundation (2015) MCI, simplified: utility = expected uses", bold=True, fmt="0.000")
    K.row("k_rec", "Material recovery", "share", "{Calc:y}", fmt=PCT)
    K.row("k_virgin", "Net virgin aluminium per use", "kg/use", "MAX(d_mass*(1-recycled_content)-{Calc:credit_mass},0)/{Calc:U}", fmt=NUM)
    K.row("k_suit", "Design suitability", "0–1",
          "IF({Calc:single}=1,1,IF({Calc:st_repair}=1,0.5+0.5*d_repair,0.5+0.5*d_disasm))", fmt="0.000")
    K.row("k_org", "Organisational reliability", "0–1", "IF({Calc:single}=1,1,0.6+0.4*bm_redeploy)", fmt="0.000")
    K.row("k_feas", "Feasibility", "0–1", "{Calc:st_feas}*{k_suit}*dist_factor*{k_org}", "base × suitability × distance × organisation", bold=True, fmt="0.000")

    K.r += 1
    K.section("Weights (edit)")
    wrows = {}
    for nm, w in WEIGHTS.items():
        K.ws.cell(row=K.r, column=1, value=f"Weight – {nm}").font = F_BASE
        c = K.ws.cell(row=K.r, column=4, value=w)
        c.font, c.border, c.number_format = F_INPUT, BOX, "0.00"
        wrows[nm] = K.r
        add_name(wb, f"w_{nm}", f"KPIs!$D${K.r}")
        K.r += 1
    K.section("Normalised criteria and weighted score")
    g, co = K.rows["k_gwp"], K.rows["k_cost"]
    K.row("n_env", "Environment (min–max, lower GWP better)", "0–1",
          f"IF(MAX($D${g}:$H${g})-MIN($D${g}:$H${g})<1E-12,0.5,1-({{k_gwp}}-MIN($D${g}:$H${g}))/(MAX($D${g}:$H${g})-MIN($D${g}:$H${g})))", fmt="0.000")
    K.row("n_circ", "Circularity", "0–1", "MIN(1,MAX(0,{k_mci}))", fmt="0.000")
    K.row("n_econ", "Economics (min–max, lower cost better)", "0–1",
          f"IF(MAX($D${co}:$H${co})-MIN($D${co}:$H${co})<1E-12,0.5,1-({{k_cost}}-MIN($D${co}:$H${co}))/(MAX($D${co}:$H${co})-MIN($D${co}:$H${co})))", fmt="0.000")
    K.row("n_feas", "Feasibility", "0–1", "MIN(1,MAX(0,{k_feas}))", fmt="0.000")
    K.row("score", "Weighted score", "0–1",
          "({n_env}*w_environment+{n_circ}*w_circularity+{n_econ}*w_economics+{n_feas}*w_feasibility)/(w_environment+w_circularity+w_economics+w_feasibility)",
          bold=True, fill=FILL_TOTAL, fmt="0.000")
    sr = K.rows["score"]
    K.row("rank", "Rank", "#", f"RANK({{score}},$D${sr}:$H${sr},0)", fmt="0", bold=True)
    K.r += 1
    K.ws.cell(row=K.r, column=1, value="Recommended strategy (highest score)").font = F_BOLD
    c = K.ws.cell(row=K.r, column=4, value=f"=INDEX($D$4:$H$4,MATCH(MAX($D${sr}:$H${sr}),$D${sr}:$H${sr},0))")
    c.font, c.fill, c.border = F_BOLD, FILL_KEY, BOX
    K.r += 1
    K.ws.cell(row=K.r, column=1, value="Lowest climate impact per use").font = F_BOLD
    c = K.ws.cell(row=K.r, column=4, value=f"=INDEX($D$4:$H$4,MATCH(MIN($D${g}:$H${g}),$D${g}:$H${g},0))")
    c.font, c.border = F_BOLD, BOX

    # ------------------------------------------------------------------ Validation
    ws = wb.create_sheet("Validation")
    title(ws, "Validation against the Python model",
          f"Reference = src/cirdesign/model.py evaluated at the default context {DEFAULT_CONTEXT['design']}/{DEFAULT_CONTEXT['business_model']}/"
          f"{DEFAULT_CONTEXT['n_uses']} uses/{DEFAULT_CONTEXT['event_km']:.0f} km/{DEFAULT_CONTEXT['credit']} when this file was built. "
          "If you change Context or Parameters the reference no longer applies.")
    ref = evaluate_all(DEFAULT_CONTEXT).join(score_strategies(DEFAULT_CONTEXT)[["score"]])
    checks = [("Expected uses", "expected_uses", "k_U"), ("Climate per use (kg CO2e)", "climate_per_use", "k_gwp"),
              ("Cost per use (EUR)", "cost_per_use", "k_cost"), ("Circularity (MCI)", "circularity", "k_mci"),
              ("Material recovery", "material_recovery", "k_rec"), ("Virgin Al per use (kg)", "virgin_al_per_use", "k_virgin"),
              ("Feasibility", "feasibility", "k_feas"), ("Weighted score", "score", "score")]
    header(ws, 4, ["KPI", "Strategy", "Python reference", "Excel", "Abs. difference", "Status"])
    r = 5
    for lab, pk, kk in checks:
        for s, col in zip(STRATS, SCOLS):
            ws.cell(row=r, column=1, value=lab)
            ws.cell(row=r, column=2, value=s)
            ws.cell(row=r, column=3, value=float(ref.loc[s, pk])).font = F_INPUT
            ws.cell(row=r, column=4, value=f"=KPIs!{col}{K.rows[kk]}").font = F_LINK
            ws.cell(row=r, column=5, value=f"=ABS(D{r}-C{r})")
            ws.cell(row=r, column=6, value=f'=IF(E{r}<=1E-6*MAX(1,ABS(C{r})),"OK","CHECK")')
            r += 1
    style_range(ws, f"A5:F{r - 1}", border=BOX)
    for row in ws[f"A5:B{r - 1}"] + ws[f"E5:F{r - 1}"]:
        for c in row:
            c.font = F_BASE
    style_range(ws, f"C5:E{r - 1}", fmt="0.000000")
    ws.cell(row=r + 1, column=1, value="Checks passing").font = F_BOLD
    ws.cell(row=r + 1, column=3, value=f'=COUNTIF(F5:F{r - 1},"OK")&" / "&COUNTA(F5:F{r - 1})').font = F_BOLD
    widths(ws, {"A": 30, "B": 14, "C": 18, "D": 18, "E": 16, "F": 10})
    ws.freeze_panes = "A5"

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = False
        for row in sh.iter_rows():
            for c in row:
                if c.font is None or c.font.name != FONT:
                    c.font = Font(name=FONT, size=c.font.size if c.font else 10, bold=c.font.b, italic=c.font.i,
                                  color=c.font.color)
    wb.active = 0
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
